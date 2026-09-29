"""엑셀 가져오기/내보내기.

가져오기: 컴시간 '교사별 시수표' 양식과 호환
    | 순 | 정식과목명 | 단축과목명 | 교사명 | 1학년 1 2 3 ... | 2학년 ... | 계 |
    - 제목 행과 학년 머리글(병합 셀) 위치는 자동으로 찾는다.
    - 셀 병합이나 순번 규칙을 엄격하게 요구하지 않는다.
    - 같은 과목·교사가 여러 줄에 나와도 합쳐서 처리한다.

내보내기: 학급별 / 교사별 / 전체 시간표 시트
"""

from __future__ import annotations

import io
import re
from collections import defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .model import Lesson, Project, SchoolClass, Subject, Teacher, new_id


class HoursTableError(ValueError):
    pass


def _cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def import_hours_table(data: bytes, project: Project, replace: bool = True) -> dict:
    """교사별 시수표 엑셀을 읽어 교사/학급/과목/수업을 만든다."""
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.worksheets[0]
    rows = [[_cell_str(c) for c in r] for r in ws.iter_rows(values_only=True)]

    # 1) '교사명' 이 있는 머리글 행 찾기
    header_idx = None
    for i, r in enumerate(rows[:30]):
        if any(x.replace(" ", "") in ("교사명", "교사", "성명") for x in r):
            header_idx = i
            break
    if header_idx is None:
        raise HoursTableError("'교사명' 머리글을 찾을 수 없습니다. 교사별 시수표 양식인지 확인해 주세요.")
    header = rows[header_idx]

    def find_col(*names):
        for j, x in enumerate(header):
            if x.replace(" ", "") in names:
                return j
        return None

    col_subject = find_col("정식과목명", "과목명", "과목")
    col_short = find_col("단축과목명", "약칭", "단축명")
    col_teacher = find_col("교사명", "교사", "성명")
    if col_subject is None:
        col_subject = col_short
    if col_subject is None or col_teacher is None:
        raise HoursTableError("과목명/교사명 열을 찾을 수 없습니다.")

    # 2) 학년/반 열 찾기: 머리글 행(학년, 병합 셀) + 다음 행(반 번호)
    grade_row = rows[header_idx]
    class_row = rows[header_idx + 1] if header_idx + 1 < len(rows) else []
    class_cols: dict[int, tuple[int, str]] = {}
    current_grade = None
    width = max(len(grade_row), len(class_row))
    for j in range(width):
        g = grade_row[j] if j < len(grade_row) else ""
        m = re.match(r"^(\d+)\s*학년$", g.replace(" ", ""))
        if m:
            current_grade = int(m.group(1))
        elif g and j > max(col_subject, col_teacher):
            if g.replace(" ", "") in ("계", "합계", "시수"):
                current_grade = None
        c = class_row[j] if j < len(class_row) else ""
        if current_grade and re.match(r"^\d+$", c) and j > max(col_subject, col_teacher):
            class_cols[j] = (current_grade, c)
    data_start = header_idx + 2
    if not class_cols:
        raise HoursTableError("학년/반 머리글을 찾을 수 없습니다. (예: '1학년' 아래에 1, 2, 3 ...)")

    if replace:
        project.teachers, project.classes, project.subjects, project.lessons, project.timetable = [], [], [], [], []

    classes: dict[tuple[int, str], SchoolClass] = {(c.grade, c.name): c for c in project.classes}
    for grade, name in class_cols.values():
        if (grade, name) not in classes:
            c = SchoolClass(id=new_id("c"), grade=grade, name=name)
            classes[(grade, name)] = c
            project.classes.append(c)
    project.classes.sort(key=lambda c: (c.grade, int(c.name) if c.name.isdigit() else 0, c.name))

    teachers: dict[str, Teacher] = {t.name: t for t in project.teachers}
    subjects: dict[str, Subject] = {s.name: s for s in project.subjects}
    hours: dict[tuple[str, str, str], int] = defaultdict(int)  # (subject, teacher, class) -> hours
    skipped = 0

    for r in rows[data_start:]:
        if not any(r):
            continue
        subj_name = r[col_subject] if col_subject < len(r) else ""
        short = r[col_short] if col_short is not None and col_short < len(r) else ""
        tname = r[col_teacher] if col_teacher < len(r) else ""
        if not subj_name and not short:
            continue
        if not tname:
            skipped += 1
            continue
        subj_name = subj_name or short
        if subj_name not in subjects:
            s = Subject(id=new_id("s"), name=subj_name, short=short or subj_name[:4])
            subjects[subj_name] = s
            project.subjects.append(s)
        if tname not in teachers:
            t = Teacher(id=new_id("t"), name=tname)
            teachers[tname] = t
            project.teachers.append(t)
        for j, key in class_cols.items():
            v = r[j] if j < len(r) else ""
            if not v:
                continue
            try:
                h = int(float(v))
            except ValueError:
                continue
            if h > 0:
                hours[(subjects[subj_name].id, teachers[tname].id, classes[key].id)] += h

    for (sid, tid, cid), h in hours.items():
        project.lessons.append(Lesson(id=new_id("l"), subject=sid, teachers=[tid], classes=[cid], hours=h))

    # 학년 교시 설정이 없으면 기본값(월~금 7교시)을 넣는다
    for grade in sorted({g for g, _ in class_cols.values()}):
        project.settings.periods.setdefault(str(grade), [7] * len(project.settings.days))

    return {"teachers": len(project.teachers), "classes": len(project.classes),
            "subjects": len(project.subjects), "lessons": len(project.lessons), "skipped_rows": skipped}


def hours_table_template() -> bytes:
    """가져오기용 빈 양식."""
    wb = Workbook()
    ws = wb.active
    ws.title = "교사별 시수표"
    ws["A1"] = "교사별 시수표"
    ws["A1"].font = Font(size=16, bold=True)
    head = ["순", "정식과목명", "단축과목명", "교사명"]
    for j, h in enumerate(head, 1):
        ws.cell(row=2, column=j, value=h)
    col = 5
    for g in (1, 2, 3):
        ws.cell(row=2, column=col, value=f"{g}학년")
        for n in range(1, 7):
            ws.cell(row=3, column=col + n - 1, value=n)
        col += 6
    ws.cell(row=2, column=col, value="계")
    example = [("국어", "국어", "홍길동", {(1, 1): 4, (1, 2): 4}), ("수학", "수학", "김수학", {(1, 1): 4})]
    for i, (s, sh, t, hs) in enumerate(example, 1):
        r = 3 + i
        ws.cell(row=r, column=1, value=i)
        ws.cell(row=r, column=2, value=s)
        ws.cell(row=r, column=3, value=sh)
        ws.cell(row=r, column=4, value=t)
        for (g, n), h in hs.items():
            ws.cell(row=r, column=5 + (g - 1) * 6 + n - 1, value=h)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------- 내보내기 ----------------

_thin = Side(style="thin", color="999999")
_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
_head_fill = PatternFill("solid", fgColor="E8EEF7")
_center = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _grid(project: Project):
    """(학급 id|교사 id, 요일, 교시) -> 표시 문자열"""
    lm, sm, tm, cm = project.lesson_map(), project.subject_map(), project.teacher_map(), project.class_map()
    by_class: dict[tuple[str, int, int], list[str]] = defaultdict(list)
    by_teacher: dict[tuple[str, int, int], list[str]] = defaultdict(list)
    for pl in project.timetable:
        l = lm.get(pl.lesson)
        if not l:
            continue
        subj = sm[l.subject].display if l.subject in sm else "?"
        tnames = ",".join(tm[t].name for t in l.teachers if t in tm)
        cnames = ",".join(cm[c].label for c in l.classes if c in cm)
        for c in l.classes:
            by_class[(c, pl.day, pl.period)].append(f"{subj}\n{tnames}")
        for t in l.teachers:
            by_teacher[(t, pl.day, pl.period)].append(f"{cnames}\n{subj}")
    return by_class, by_teacher


def _write_block(ws, top: int, title: str, owner: str, grid, project: Project, n_periods: int) -> int:
    s = project.settings
    days = s.days
    ws.cell(row=top, column=1, value=title).font = Font(bold=True, size=12)
    ws.cell(row=top + 1, column=1, value="교시").fill = _head_fill
    for d, name in enumerate(days):
        c = ws.cell(row=top + 1, column=2 + d, value=name)
        c.fill, c.alignment, c.border = _head_fill, _center, _border
    ws.cell(row=top + 1, column=1).border = _border
    for p in range(n_periods):
        label = f"{p + 1}"
        if p < len(s.bell) and s.bell[p]:
            label += f"\n{s.bell[p]}"
        c = ws.cell(row=top + 2 + p, column=1, value=label)
        c.alignment, c.border, c.fill = _center, _border, _head_fill
        for d in range(len(days)):
            v = " / ".join(grid.get((owner, d, p), []))
            c = ws.cell(row=top + 2 + p, column=2 + d, value=v)
            c.alignment, c.border = _center, _border
        ws.row_dimensions[top + 2 + p].height = 32
    return top + n_periods + 3


def export_timetable(project: Project) -> bytes:
    s = project.settings
    n_periods = s.max_periods
    by_class, by_teacher = _grid(project)
    wb = Workbook()

    # 전체 시간표 (학급 x 요일·교시)
    ws = wb.active
    ws.title = "전체(학급)"
    ws.cell(row=1, column=1, value=f"{s.school_name} {s.year}학년도 {s.semester}학기 전체 시간표").font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value="학급").fill = _head_fill
    col = 2
    for d, name in enumerate(s.days):
        for p in range(n_periods):
            c = ws.cell(row=2, column=col, value=f"{name}{p + 1}")
            c.fill, c.alignment, c.border = _head_fill, _center, _border
            ws.column_dimensions[get_column_letter(col)].width = 7
            col += 1
    for i, cls in enumerate(project.classes):
        ws.cell(row=3 + i, column=1, value=cls.label).border = _border
        col = 2
        for d in range(len(s.days)):
            for p in range(n_periods):
                v = by_class.get((cls.id, d, p), [])
                c = ws.cell(row=3 + i, column=col, value=" / ".join(x.split("\n")[0] for x in v))
                c.alignment, c.border = _center, _border
                col += 1
    ws.freeze_panes = "B3"

    ws = wb.create_sheet("학급별")
    ws.column_dimensions["A"].width = 12
    for d in range(len(s.days)):
        ws.column_dimensions[get_column_letter(2 + d)].width = 14
    row = 1
    tm = project.teacher_map()
    for cls in project.classes:
        title = f"{cls.label} 시간표" + (f" (담임 {tm[cls.homeroom].name})" if cls.homeroom in tm else "")
        row = _write_block(ws, row, title, cls.id, by_class, project, n_periods)

    ws = wb.create_sheet("교사별")
    ws.column_dimensions["A"].width = 12
    for d in range(len(s.days)):
        ws.column_dimensions[get_column_letter(2 + d)].width = 14
    row = 1
    for t in project.teachers:
        n = sum(1 for (tid, _, _) in by_teacher if tid == t.id)
        row = _write_block(ws, row, f"{t.name} 선생님 ({n}시간)", t.id, by_teacher, project, n_periods)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
