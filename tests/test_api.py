import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook


@pytest.fixture()
def client(tmp_path, monkeypatch):
    from app import main
    from app.storage import Storage

    monkeypatch.setattr(main, "store", Storage(tmp_path))
    return TestClient(main.app)


def test_index(client):
    r = client.get("/")
    assert r.status_code == 200
    assert "시간표" in r.text


def test_project_crud_and_backup(client):
    p = client.post("/api/projects", json={"title": "테스트"}).json()
    assert p["title"] == "테스트"
    p["title"] = "바뀐 이름"
    assert client.put(f"/api/projects/{p['id']}", json=p).status_code == 200
    assert client.get(f"/api/projects/{p['id']}").json()["title"] == "바뀐 이름"
    backups = client.get(f"/api/projects/{p['id']}/backups").json()
    assert len(backups) == 1
    restored = client.post(f"/api/projects/{p['id']}/backups/{backups[0]['name']}/restore").json()
    assert restored["title"] == "테스트"
    assert [x["id"] for x in client.get("/api/projects").json()] == [p["id"]]
    client.delete(f"/api/projects/{p['id']}")
    assert client.get(f"/api/projects/{p['id']}").status_code == 404


def test_bad_project_id_rejected(client):
    assert client.get("/api/projects/..%2F..%2Fetc").status_code == 404


def test_sample_check_solve_export(client):
    p = client.post("/api/projects", json={"sample": True}).json()
    chk = client.post("/api/check", json=p).json()
    assert chk["summary"]["data"]["error"] == 0
    r = client.post(f"/api/projects/{p['id']}/solve", json={"options": {"time_limit": 60}}).json()
    assert r["status"] in ("optimal", "feasible"), r
    saved = client.get(f"/api/projects/{p['id']}").json()
    assert len(saved["timetable"]) == len(r["timetable"]) > 0
    chk = client.post("/api/check", json=saved).json()
    assert chk["summary"]["timetable"]["error"] == 0
    x = client.get(f"/api/projects/{p['id']}/export.xlsx")
    wb = load_workbook(io.BytesIO(x.content))
    assert {"전체(학급)", "학급별", "교사별"} <= set(wb.sheetnames)


def _hours_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "교사별 시수표"
    for j, h in enumerate(["순", "정식과목명", "단축과목명", "교사명"], 1):
        ws.cell(row=2, column=j, value=h)
    ws.cell(row=2, column=5, value="1학년")
    ws.cell(row=2, column=7, value="2학년")
    ws.cell(row=2, column=9, value="계")
    for j, n in enumerate([1, 2, 1, 2], 5):
        ws.cell(row=3, column=j, value=n)
    data = [
        (1, "국어", "국어", "홍길동", [4, 4, None, None]),
        (2, "수학", "수학", "김수학", [4, None, 4, 4]),
        (3, "국어", "국어", "홍길동", [None, None, 2, None]),  # 같은 교사·과목이 여러 줄
    ]
    for i, (n, s, sh, t, hs) in enumerate(data):
        r = 4 + i
        for j, v in enumerate([n, s, sh, t], 1):
            ws.cell(row=r, column=j, value=v)
        for j, h in enumerate(hs, 5):
            if h:
                ws.cell(row=r, column=j, value=h)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_import_hours_table(client):
    p = client.post("/api/projects", json={"title": "가져오기"}).json()
    r = client.post(f"/api/projects/{p['id']}/import-hours", content=_hours_xlsx(),
                    headers={"Content-Type": "application/octet-stream"})
    assert r.status_code == 200, r.text
    stats = r.json()["stats"]
    # 국어(홍길동) 1-1, 1-2, 2-1 / 수학(김수학) 1-1, 2-1, 2-2
    assert stats == {"teachers": 2, "classes": 4, "subjects": 2, "lessons": 6, "skipped_rows": 0}
    proj = r.json()["project"]
    labels = sorted(f"{c['grade']}-{c['name']}" for c in proj["classes"])
    assert labels == ["1-1", "1-2", "2-1", "2-2"]


def test_import_rejects_wrong_file(client):
    p = client.post("/api/projects", json={"title": "x"}).json()
    r = client.post(f"/api/projects/{p['id']}/import-hours", content=b"not excel",
                    headers={"Content-Type": "application/octet-stream"})
    assert r.status_code == 400


def test_template_download(client):
    r = client.get("/api/template.xlsx")
    assert r.status_code == 200
    load_workbook(io.BytesIO(r.content))
